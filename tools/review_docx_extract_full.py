from docx import Document

path = r"D:\fileless malware\mã độc.docx"
out = r"D:\fileless malware\ma_doc_extracted_review.txt"

doc = Document(path)

with open(out, "w", encoding="utf-8") as handle:
    handle.write("=== PARAGRAPHS ===\n")
    for i, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            handle.write(f"{i:04d} | {paragraph.style.name} | {text}\n")

    handle.write("\n=== TABLES ===\n")
    for t_i, table in enumerate(doc.tables):
        handle.write(f"\n--- TABLE {t_i} ---\n")
        for r_i, row in enumerate(table.rows):
            cells = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
            handle.write(f"{r_i:03d}: " + " | ".join(cells) + "\n")

print(out)
