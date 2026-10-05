from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


DOCX_PATH = r"D:\fileless malware\Fileless_Malware.docx"


def set_run_font(run, size=11, bold=False, color="111111", italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", style=None, bold=False, italic=False):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text)
        set_run_font(r, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    style = f"Heading {level}" if level in (1, 2, 3) else None
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r, size=14 if level == 1 else 12, bold=True, color="000000")
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-10)
        r = p.add_run("• ")
        set_run_font(r)
        r = p.add_run(item)
        set_run_font(r)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{margin}"))
        if element is None:
            element = OxmlElement(f"w:{margin}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "1F2937")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=10.5, bold=True, color="FFFFFF")
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cell = cells[i]
            if row_index % 2 == 1:
                set_cell_shading(cell, "F7F7F7")
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_run_font(r, size=10)
    doc.add_paragraph()
    return table


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F2F2F2")
    set_cell_border(cell, "D9D9D9")
    set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(lines):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        r.font.size = Pt(9.5)
    doc.add_paragraph()


doc = Document(DOCX_PATH)
doc.add_page_break()

add_heading(doc, "CHƯƠNG 3: DEMO LAB MÔ PHỎNG ASTAROTH-INSPIRED FILELESS MALWARE", 1)

add_heading(doc, "3.1. Mục tiêu và lý do chọn Astaroth làm bài lab", 2)
add_para(
    doc,
    "Astaroth được chọn làm đối tượng mô phỏng trong bài lab vì đây là một trường hợp tiêu biểu cho xu hướng Fileless Malware hiện đại: không phụ thuộc vào một tệp thực thi độc hại đơn lẻ, mà tận dụng nhiều công cụ hợp pháp có sẵn trong Windows để tải, giải mã và kích hoạt payload. Trong phần lý thuyết, Astaroth thể hiện rõ các đặc điểm quan trọng như tấn công qua phishing, sử dụng file nén chứa shortcut LNK, lạm dụng các công cụ hệ thống và che giấu payload qua nhiều lớp trung gian.",
)
add_para(
    doc,
    "Tuy nhiên, lab không sử dụng mẫu mã độc thật và không tái hiện các bước nguy hiểm như đánh cắp thông tin, ADS, nạp DLL độc hại, process injection hoặc kết nối C2 thật. Thay vào đó, lab chỉ mô phỏng rút gọn các ý tưởng chính của Astaroth bằng payload vô hại, nhằm phục vụ mục tiêu học tập và phân tích log.",
)
add_bullets(
    doc,
    [
        "Astaroth có chuỗi tấn công phù hợp để minh họa kỹ thuật Living off the Land.",
        "Chuỗi tấn công có yếu tố social engineering, giúp lab có kịch bản gần thực tế hơn.",
        "Các bước như LNK, BITS/PowerShell download và Certutil decode đều tạo log rõ ràng để phân tích.",
        "Lab có thể triển khai an toàn trong môi trường hai máy mà không cần dùng sample mã độc thật.",
    ],
)

add_heading(doc, "3.2. Phạm vi mô phỏng và giới hạn an toàn", 2)
add_para(
    doc,
    "Để đảm bảo an toàn, bài lab được thiết kế theo hướng Astaroth-inspired, tức là lấy cảm hứng từ chuỗi kỹ thuật của Astaroth nhưng không sao chép đầy đủ hành vi mã độc. Lab tập trung vào các dấu hiệu có thể quan sát bằng log, đặc biệt là chuỗi process, command line, network connection và hành vi decode payload.",
)
add_table(
    doc,
    ["Thành phần trong Astaroth thật", "Cách mô phỏng trong lab", "Ghi chú an toàn"],
    [
        ["Email phishing hoặc link độc hại", "Trang web giả lập nội bộ trên máy attacker", "Không gửi email thật, không thu thập thông tin người dùng"],
        ["File ZIP chứa shortcut LNK", "ZIP chứa file ThongBaoLuong_Q3.lnk", "Người dùng chủ động mở trong môi trường lab"],
        ["Chuỗi script nhiều lớp", "Shortcut gọi stage XacNhanLuong_Q3.cmd", "Stage chỉ tải payload vô hại"],
        ["BITSAdmin tải payload", "PowerShell/BITS tải payload.b64", "Payload đặt trong mạng lab nội bộ"],
        ["Certutil decode payload", "certutil.exe decode payload.b64 thành payload-decoded.ps1", "Payload sau decode không có hành vi độc hại"],
        ["ADS, ExtExport, DLL injection, credential theft", "Không mô phỏng", "Loại bỏ để tránh triển khai hành vi nguy hiểm"],
    ],
)

add_heading(doc, "3.3. Môi trường lab", 2)
add_para(doc, "Lab được triển khai trên mô hình hai máy để thể hiện rõ vai trò attacker và nạn nhân.")
add_table(
    doc,
    ["Thành phần", "Vai trò", "Địa chỉ IP"],
    [
        ["Máy thật Windows", "Máy attacker, host trang web giả lập và payload vô hại", "192.168.141.1"],
        ["Máy ảo Windows", "Máy nạn nhân, truy cập link, tải ZIP và sinh log phân tích", "192.168.141.140"],
    ],
)
add_para(doc, "Trên máy nạn nhân cần bật các nguồn log sau để phục vụ phân tích:")
add_bullets(
    doc,
    [
        "Sysmon Event ID 1 để ghi nhận process creation.",
        "Sysmon Event ID 3 để ghi nhận network connection.",
        "Sysmon Event ID 11 để ghi nhận file create.",
        "PowerShell Event ID 4104 để ghi nhận script block.",
        "Security Event ID 4688 để ghi nhận process creation kèm command line.",
    ],
)

add_heading(doc, "3.4. Kịch bản tổng quan của bài lab", 2)
add_para(
    doc,
    "Máy attacker host một trang web giả lập cổng tài liệu nội bộ. Nạn nhân truy cập trang này, tải file ZIP giả dạng hồ sơ xác nhận lương quý 3, giải nén và mở shortcut LNK bên trong. Shortcut đóng vai trò tương tự bước kích hoạt ban đầu trong chuỗi Astaroth, sau đó tải stage trung gian từ máy attacker. Stage này tiếp tục tải payload vô hại, decode bằng Certutil và chạy bằng PowerShell.",
)
add_code_block(
    doc,
    [
        "Attacker host trang web giả lập",
        "  -> Nạn nhân tải ThongBaoLuong_Q3.zip",
        "  -> Nạn nhân giải nén và mở ThongBaoLuong_Q3.lnk",
        "  -> Shortcut gọi cmd.exe / PowerShell",
        "  -> Tải stage XacNhanLuong_Q3.cmd từ 192.168.141.1",
        "  -> Stage tải payload.b64",
        "  -> certutil decode payload.b64",
        "  -> PowerShell chạy payload-decoded.ps1",
        "  -> Popup giả lập xuất hiện và lab-result.txt được tạo",
    ],
)

add_heading(doc, "3.5. Các bước thực hiện lab", 2)
add_heading(doc, "3.5.1. Chuẩn bị trên máy attacker", 3)
add_para(doc, "Trên máy attacker, tạo thư mục web server và copy các file lab vào thư mục này.")
add_code_block(
    doc,
    [
        "C:\\Lab\\Fileless-Attacker\\web",
        "index.html",
        "payload.ps1",
        "XacNhanLuong_Q3.cmd",
        "create_lnk_zip.ps1",
    ],
)
add_para(doc, "Sau đó tạo payload Base64 và tạo file ZIP chứa shortcut LNK. Cuối cùng bật HTTP server trên địa chỉ 192.168.141.1 cổng 8000.")
add_code_block(
    doc,
    [
        "certutil -encode payload.ps1 payload.b64",
        "PowerShell -ExecutionPolicy Bypass -File create_lnk_zip.ps1",
        "python -m http.server 8000 --bind 192.168.141.1",
    ],
)

add_heading(doc, "3.5.2. Thao tác trên máy nạn nhân", 3)
add_bullets(
    doc,
    [
        "Truy cập trang giả lập tại http://192.168.141.1:8000/.",
        "Tải file ThongBaoLuong_Q3.zip.",
        "Giải nén file ZIP và mở shortcut ThongBaoLuong_Q3.lnk.",
        "Quan sát popup giả lập xuất hiện sau khi payload vô hại được chạy.",
        "Kiểm tra thư mục C:\\Lab\\Fileless\\work để xác nhận các file payload.b64, payload-decoded.ps1 và lab-result.txt được tạo.",
    ],
)

add_heading(doc, "3.5.3. Dữ liệu cần chụp và thu thập", 3)
add_bullets(
    doc,
    [
        "Cửa sổ HTTP server trên máy attacker có request từ IP 192.168.141.140.",
        "Request tải ThongBaoLuong_Q3.zip, XacNhanLuong_Q3.cmd và payload.b64.",
        "File ZIP và shortcut LNK trên máy nạn nhân.",
        "Popup giả lập sau khi payload chạy thành công.",
        "Sysmon Event ID 1 cho cmd.exe, powershell.exe và certutil.exe.",
        "Sysmon Event ID 3 cho kết nối đến 192.168.141.1:8000.",
        "PowerShell Event ID 4104 ghi nhận nội dung script block.",
    ],
)

add_heading(doc, "3.6. Kết quả rút ra từ bài lab", 2)
add_para(
    doc,
    "Sau khi thực hiện lab, có thể thấy rằng Fileless Malware không nhất thiết cần triển khai một file thực thi độc hại cố định trên ổ đĩa. Chuỗi tấn công có thể bắt đầu từ một file ZIP chứa shortcut LNK, sau đó chuyển sang lạm dụng các công cụ hợp pháp của Windows để tải và xử lý payload. Trong lab, payload hoàn toàn vô hại, nhưng chuỗi hành vi vẫn tạo ra các dấu hiệu quan trọng để phân tích phòng thủ.",
)
add_bullets(
    doc,
    [
        "Shortcut LNK là điểm khởi phát đáng chú ý vì nó có thể kích hoạt dòng lệnh mà người dùng không dễ nhận ra nếu chỉ nhìn tên file.",
        "PowerShell/BITS tạo dấu hiệu network connection từ máy nạn nhân đến máy attacker.",
        "Certutil decode payload là hành vi đáng giám sát vì công cụ này thường bị lạm dụng trong các chuỗi Living off the Land.",
        "Các sự kiện riêng lẻ có thể chưa đủ kết luận độc hại, nhưng khi ghép thành chuỗi theo thời gian, chúng tạo thành dấu hiệu tấn công rõ ràng.",
        "Lab cho thấy phát hiện Fileless Malware cần dựa vào behavior-based detection, process chain, command line và correlation nhiều nguồn log.",
    ],
)

add_heading(doc, "3.7. Hạn chế của bài lab", 2)
add_para(
    doc,
    "Bài lab chỉ mô phỏng rút gọn để phục vụ học tập nên không phản ánh toàn bộ độ phức tạp của Astaroth thật. Các kỹ thuật nguy hiểm như ADS, ExtExport, nạp DLL, tiêm tiến trình, đánh cắp thông tin và kết nối C2 thật không được triển khai. Đây là giới hạn có chủ đích để đảm bảo lab an toàn, dễ kiểm soát và phù hợp với mục tiêu phân tích phòng thủ.",
)

add_heading(doc, "3.8. Kết luận chương", 2)
add_para(
    doc,
    "Thông qua lab, có thể thấy lý do Astaroth là một case phù hợp để nghiên cứu Fileless Malware: chuỗi tấn công kết hợp social engineering, shortcut LNK, công cụ hệ thống hợp pháp và payload nhiều lớp. Bài lab đã mô phỏng các yếu tố cốt lõi ở mức an toàn, giúp người học quan sát được quá trình tải payload, decode payload, thực thi payload vô hại và phân tích log tương ứng. Kết quả quan trọng nhất là việc phòng thủ trước Fileless Malware cần tập trung vào hành vi tổng thể thay vì chỉ tìm kiếm một file độc hại riêng lẻ.",
)

doc.save(DOCX_PATH)
print(DOCX_PATH)
