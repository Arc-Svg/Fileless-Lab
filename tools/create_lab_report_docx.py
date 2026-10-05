from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"D:\fileless malware\bao-cao-ngan-lab-fileless-malware.docx"


def set_run_font(run, size=10.5, bold=False, color="111827"):
    run.font.name = "Segoe UI"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Segoe UI")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Segoe UI")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{margin}"))
        if element is None:
            element = OxmlElement(f"w:{margin}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def add_paragraph(doc, text="", style=None, size=10.5, bold=False, color="111827", align=None):
    paragraph = doc.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.12
    if text:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True, color="000000")
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        set_run_font(run)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        set_run_font(run)


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F3F4F6")
    set_cell_border(cell, "D1D5DB")
    set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(lines):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string("111827")
    doc.add_paragraph()


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        cell = header_cells[i]
        set_cell_shading(cell, "1F2937")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_run_font(run, bold=True, color="FFFFFF")
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cell = cells[i]
            if row_index % 2 == 1:
                set_cell_shading(cell, "F9FAFB")
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            set_run_font(run, size=10)
    doc.add_paragraph()
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
for name in ("Normal", "Title", "Heading 1", "Heading 2"):
    styles[name].font.name = "Segoe UI"
    styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Segoe UI")
    styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Segoe UI")
styles["Normal"].font.size = Pt(10.5)
styles["Title"].font.size = Pt(22)
styles["Title"].font.bold = True

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Báo cáo ngắn đề xuất lab Fileless Malware")
set_run_font(run, size=22, bold=True, color="000000")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(14)
run = subtitle.add_run("Mô phỏng an toàn chuỗi Astaroth inspired Living off the Land")
set_run_font(run, size=11, color="4B5563")

add_heading(doc, "1. Tên lab")
add_paragraph(doc, "Mô phỏng an toàn chuỗi tấn công Fileless Malware theo hướng Astaroth inspired Living off the Land.")

add_heading(doc, "2. Mục đích của lab")
add_paragraph(
    doc,
    "Lab được xây dựng nhằm mô phỏng một kịch bản Fileless Malware ở mức an toàn, phục vụ cho mục tiêu học tập và phân tích phòng thủ. Thay vì sử dụng mẫu mã độc thật, lab chỉ dùng payload vô hại để tạo file kết quả và hiển thị popup giả lập.",
)
add_bullets(
    doc,
    [
        "Minh họa cách kẻ tấn công có thể lạm dụng các công cụ hợp pháp có sẵn trong Windows.",
        "Quan sát chuỗi hành vi từ lúc nạn nhân truy cập trang giả lập đến khi payload vô hại được thực thi.",
        "Thu thập log trên máy nạn nhân bằng Sysmon, PowerShell Logging và Event Viewer.",
        "Phân tích dấu hiệu phát hiện dựa trên process chain, command line, network connection và hành vi decode payload.",
        "Rút ra cách phòng thủ trước các kỹ thuật Living off the Land thường gặp trong Fileless Malware.",
    ],
)
add_paragraph(
    doc,
    "Lab lấy cảm hứng từ Astaroth vì đây là case tiêu biểu được Microsoft phân tích trong nhóm fileless và living off the land. Tuy nhiên, lab không tái hiện đầy đủ chuỗi Astaroth thật, mà chỉ mô phỏng rút gọn một số ý tưởng cốt lõi để đảm bảo an toàn.",
)

add_heading(doc, "3. Môi trường chuẩn bị")
add_table(
    doc,
    ["Thành phần", "Vai trò", "Địa chỉ IP"],
    [
        ["Máy thật Windows", "Máy attacker, host trang web giả lập và payload vô hại", "192.168.141.1"],
        ["Máy ảo Windows", "Máy nạn nhân, truy cập link và thực hiện thao tác giả lập", "192.168.141.140"],
    ],
)
add_paragraph(doc, "Các công cụ cần chuẩn bị trên máy nạn nhân:")
add_bullets(doc, ["Windows Event Viewer.", "Sysmon.", "PowerShell Logging.", "Security Event ID 4688 kèm command line.", "Trình duyệt web."])
add_paragraph(doc, "Các công cụ cần chuẩn bị trên máy attacker:")
add_bullets(doc, ["Python HTTP server.", "Thư mục web chứa trang giả lập, file tài liệu mồi và payload vô hại."])

add_heading(doc, "4. Kịch bản mô phỏng")
add_paragraph(
    doc,
    "Kịch bản lab mô phỏng tình huống thường gặp trong thực tế: nạn nhân nhận được một đường link giả lập dẫn đến trang tài liệu nội bộ. Trang này yêu cầu tải tài liệu xác nhận thông tin lương.",
)
add_code_block(
    doc,
    [
        "Attacker host trang web giả lập",
        "  -> Nạn nhân truy cập link nội bộ giả",
        "  -> Nạn nhân tải file Word ThongBaoLuong_Q3.docx",
        "  -> Nạn nhân mở tài liệu và bấm link xác nhận",
        "  -> File xác nhận gọi PowerShell/BITS tải payload vô hại",
        "  -> certutil decode payload",
        "  -> PowerShell chạy payload vô hại",
        "  -> Popup giả lập xuất hiện và file kết quả được tạo",
    ],
)
add_paragraph(doc, "Payload cuối cùng không có chức năng độc hại. Nó chỉ tạo file lab-result.txt trên máy nạn nhân và hiển thị popup giả lập để cho thấy payload đã chạy thành công.")

add_heading(doc, "5. Các bước thực hiện cơ bản")
add_numbered(
    doc,
    [
        "Trên máy attacker, tạo thư mục C:\\Lab\\Fileless-Attacker\\web và copy bộ file lab từ D:\\fileless malware\\lab-files\\attacker-web.",
        "Encode payload vô hại bằng certutil để tạo file payload.b64.",
        "Mở HTTP server bằng lệnh python -m http.server 8000 --bind 192.168.141.1.",
        "Trên máy nạn nhân, kiểm tra kết nối đến máy attacker bằng ping 192.168.141.1.",
        "Truy cập trang giả lập tại http://192.168.141.1:8000/.",
        "Tải và mở file Word ThongBaoLuong_Q3.docx.",
        "Bấm liên kết Mở biểu mẫu xác nhận trong tài liệu.",
        "Quan sát payload vô hại được thực thi, popup xuất hiện và file lab-result.txt được tạo.",
    ],
)

add_heading(doc, "6. Dữ liệu và log cần thu thập")
add_table(
    doc,
    ["Nguồn log", "Nội dung cần quan sát"],
    [
        ["Máy attacker", "HTTP request từ 192.168.141.140, gồm request tải tài liệu, file xác nhận và payload.b64."],
        ["Sysmon Event ID 1", "Process Creation liên quan cmd.exe, powershell.exe và certutil.exe."],
        ["Sysmon Event ID 3", "Network Connection từ máy nạn nhân đến 192.168.141.1:8000."],
        ["Sysmon Event ID 11", "File Create cho payload.b64, payload-decoded.ps1 và lab-result.txt."],
        ["PowerShell Event ID 4104", "Script Block Logging cho lệnh tải payload và chạy script."],
        ["Security Event ID 4688", "Process Creation kèm command line để phân tích parent-child process."],
    ],
)
add_paragraph(doc, "Các dấu hiệu đáng chú ý gồm PowerShell tải file từ IP nội bộ lạ, certutil.exe được dùng để decode file, và chuỗi tải payload, decode payload, chạy script xuất hiện trong thời gian ngắn.")

add_heading(doc, "7. Kết quả dự kiến")
add_bullets(
    doc,
    [
        "Trang web giả lập đóng vai trò mồi nhử để nạn nhân tải tài liệu.",
        "File tài liệu Word dẫn nạn nhân đến bước thực thi tiếp theo thông qua social engineering.",
        "Các công cụ hợp pháp như PowerShell/BITS và certutil có thể bị lạm dụng để tải và xử lý payload.",
        "Dù không dùng file mã độc thật, hệ thống vẫn sinh ra nhiều log có giá trị để phát hiện hành vi đáng ngờ.",
        "Việc phát hiện Fileless Malware nên dựa vào hành vi và tương quan sự kiện, không chỉ dựa vào chữ ký file.",
    ],
)

add_heading(doc, "8. Kết luận rút ra")
add_paragraph(
    doc,
    "Lab cho thấy Fileless Malware không nhất thiết phải thả một file .exe độc hại rõ ràng xuống ổ đĩa. Thay vào đó, kẻ tấn công có thể lợi dụng công cụ hợp pháp của Windows để thực hiện các bước như tải payload, decode payload và chạy script.",
)
add_code_block(
    doc,
    [
        "Người dùng mở tài liệu đáng ngờ",
        "  -> cmd.exe hoặc script được gọi",
        "  -> PowerShell/BITS tải file",
        "  -> certutil decode nội dung",
        "  -> PowerShell chạy payload",
    ],
)
add_paragraph(
    doc,
    "Từng hành vi riêng lẻ có thể có lý do hợp pháp, nhưng khi xuất hiện liên tiếp trong thời gian ngắn thì trở thành dấu hiệu đáng nghi. Vì vậy, phòng thủ nên tập trung vào logging đầy đủ, giám sát LOLBins, theo dõi parent-child process bất thường, phân tích command line, áp dụng AppLocker hoặc WDAC và đào tạo người dùng về tài liệu hoặc link giả mạo.",
)

doc.save(OUTPUT)
print(OUTPUT)
